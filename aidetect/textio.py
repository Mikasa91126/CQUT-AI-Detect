"""文本输入：txt / md / docx / pdf / stdin，统一成纯文本。

约定：任何读取失败都给中文可读的报错，不抛裸 traceback 给终端用户。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".text", ".log", ".csv"}
DOCX_SUFFIXES = {".docx"}
PDF_SUFFIXES = {".pdf"}

# 二进制嗅探：NUL 字节或过多控制字符即判为「不是文本」
_BINARY_SNIFF_BYTES = 4096
_CONTROL_RATIO_MAX = 0.05
_UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")
# 终端彩色日志里的 ANSI 转义序列（ESC [ ... m 等）不算「二进制特征」
_ANSI_ESCAPE = re.compile(rb"\x1b\[[0-9;?]*[ -/]*[@-~]")

BINARY_HINT = "无法解析文档：这看起来是二进制文件，不是文本（支持 .txt / .md / .docx / .pdf）"


class ReadError(RuntimeError):
    """输入不可用（文件缺失、格式不支持、内容为空）。"""


def _looks_binary(data: bytes) -> bool:
    """判断字节流是否更像二进制文件而非文本（UTF-16 BOM 视为文本）。"""
    if not data:
        return False
    if data.startswith(_UTF16_BOMS):
        return False
    sample = data[:_BINARY_SNIFF_BYTES]
    if b"\x00" in sample:
        return True
    sample = _ANSI_ESCAPE.sub(b"", sample)
    if not sample:
        return False
    ctrl = sum(1 for b in sample if b < 0x09 or 0x0E <= b < 0x20)
    return ctrl / len(sample) > _CONTROL_RATIO_MAX


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


# 解码后残留的 ANSI 转义序列与控制字符：清掉，避免被当成「正文」计分
_ANSI_TEXT = re.compile(r"\x1b\[[0-9;?]*[ -/]*[@-~]")
_CTRL_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean_text(text: str) -> str:
    return _CTRL_TEXT.sub("", _ANSI_TEXT.sub("", text))


def _read_plain_text(path: Path, hint: str = "") -> str:
    data = path.read_bytes()
    if _looks_binary(data):
        raise ReadError(BINARY_HINT + hint)
    text = _clean_text(_decode(data))
    if not text.strip():
        raise ReadError("文件没有可读文本（只含控制字符或 ANSI 转义序列）")
    return text


def read_text_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in DOCX_SUFFIXES:
        return _read_docx(path)
    if suffix in PDF_SUFFIXES:
        return _read_pdf(path)
    if suffix in TEXT_SUFFIXES or suffix == "":
        return _read_plain_text(path)
    # 未知后缀：先当纯文本试一次，再当二进制失败
    return _read_plain_text(path, hint=f"（后缀 {suffix or '无'}）")


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - 环境缺依赖
        raise ReadError("缺少 python-docx，无法读取 .docx（pip install python-docx）") from exc
    try:
        d = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"docx 打开失败（可能不是标准 docx）: {exc}") from exc
    blocks: list[str] = []
    for p in d.paragraphs:
        t = (p.text or "").strip()
        if t:
            blocks.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            line = " | ".join(x for x in cells if x)
            if line:
                blocks.append(line)
    return "\n\n".join(blocks)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        PdfReader = None  # type: ignore[assignment]
    if PdfReader is not None:
        try:
            reader = PdfReader(str(path))
            pages = [(pg.extract_text() or "") for pg in reader.pages]
            text = "\n".join(pages).strip()
            if text:
                return text
        except Exception:  # noqa: BLE001
            pass
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise ReadError("PDF 读取失败：pypdf 解析不出文本，且未安装 PyMuPDF") from exc
    try:
        with fitz.open(str(path)) as doc:
            return "\n".join(page.get_text() for page in doc).strip()
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"PDF 读取失败: {exc}") from exc


def read_stdin() -> str:
    data = sys.stdin.buffer.read()
    if not data:
        raise ReadError("stdin 没有读到内容")
    text = _clean_text(_decode(data))
    if not text.strip():
        raise ReadError("stdin 没有可读文本（只含控制字符或 ANSI 转义序列）")
    return text


def load(source: str) -> tuple[str, str]:
    """返回 (文本, 展示名)。source 为 '-' 时读 stdin。"""
    if source == "-":
        return read_stdin(), "<stdin>"
    path = Path(source).expanduser()
    if not path.exists():
        raise ReadError(f"文件不存在: {path}")
    if path.is_dir():
        raise ReadError(f"这是一个目录，不是文件: {path}")
    text = read_text_file(path)
    if not text.strip():
        raise ReadError(f"文件没有可读文本: {path}")
    return text, path.name
